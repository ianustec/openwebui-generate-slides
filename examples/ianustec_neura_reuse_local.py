"""Local battery for the NEURA reuse JSON + Marketing.pptx.

Saves via Tools._save cache fallback into the repo (not /app/backend/...):

    examples/cache/files/

    python examples/ianustec_neura_reuse_local.py
"""
from __future__ import annotations

import asyncio
import importlib.util
import json
import sys
import zipfile
import re
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

BASE = Path(__file__).resolve().parent.parent
SPEC_PATH = BASE / "examples" / "ianustec_neura_reuse.json"
MARKETING = BASE / "doc" / "online_templates" / "Marketing.pptx"
CACHE_DIR = BASE / "examples" / "cache" / "files"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _slide_stats(pptx_bytes: bytes, index: int) -> dict:
    z = zipfile.ZipFile(BytesIO(pptx_bytes))
    xml = z.read(f"ppt/slides/slide{index + 1}.xml").decode("utf-8", "replace")
    texts = re.findall(r"<a:t>([^<]*)</a:t>", xml)
    return {
        "sp": len(re.findall(r"<p:sp\b", xml)),
        "grp": len(re.findall(r"<p:grpSp\b", xml)),
        "pic": len(re.findall(r"<p:pic\b", xml)),
        "cust": len(re.findall(r"<a:custGeom", xml)),
        "blip": len(re.findall(r"<a:blip\b", xml)),
        "blip_embed": len(re.findall(r'r:embed="[^"]+"', xml)),
        "svg": len(re.findall(r"<asvg:svgBlip\b", xml)),
        # distinct letter-spacing values: text replacement merges adjacent runs
        # with equal formatting, so run *count* is not a stable metric.
        "spc": sorted(set(re.findall(r'\bspc="(-?\d+)"', xml)), key=int),
        "texts": texts[:12],
        "dangling": _dangling_embeds(z, index, xml),
    }


def _dangling_embeds(z: zipfile.ZipFile, index: int, xml: str) -> list[str]:
    """r:embed ids used by the slide XML that have no relationship target."""
    rels_name = f"ppt/slides/_rels/slide{index + 1}.xml.rels"
    rels_xml = z.read(rels_name).decode("utf-8", "replace") if rels_name in z.namelist() else ""
    known = set(re.findall(r'Id="(rId\d+)"', rels_xml))
    used = set(re.findall(r'r:embed="(rId\d+)"', xml))
    return sorted(used - known)


def _shape_text(slide, shape_id: int, mod) -> str | None:
    sh = mod._shape_by_id(slide, shape_id)
    if sh is None or not getattr(sh, "has_text_frame", False):
        return None
    return (sh.text or "").strip()


async def _save_via_cache(mod, data: bytes, title: str) -> Path:
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    tools = mod.Tools()
    tools.valves.pptx_export_dir = str(CACHE_DIR)
    fname, url, err, file_id = await tools._save(
        data, title=title, request=None, user_dict=None
    )
    if err or not fname:
        raise AssertionError(f"_save cache fallback failed: {err}")
    path = CACHE_DIR / fname
    if not path.is_file():
        raise AssertionError(f"expected cache file {path}")
    print(f"saved {path} ({path.stat().st_size} bytes) url={url} file_id={file_id}")
    return path


def main() -> None:
    if not MARKETING.is_file():
        print(f"MISSING {MARKETING}", file=sys.stderr)
        sys.exit(2)
    if not SPEC_PATH.is_file():
        print(f"MISSING {SPEC_PATH}", file=sys.stderr)
        sys.exit(2)

    mod = _load_mod()
    spec = json.loads(SPEC_PATH.read_text())
    ref = MARKETING.read_bytes()
    pack = mod._parse_reference_pptx(ref)

    print("=== parse Marketing ===")
    print(
        f"slides={len(pack.slides)} size={pack.slide_width_in}x{pack.slide_height_in}"
    )

    print("=== _build (reuse path, no Files API) ===")
    data, n = mod.Tools()._build(
        spec, template_pack=pack, reference_bytes=ref
    )
    print(f"build slides={n} bytes={len(data)}")
    if n != 11:
        raise AssertionError(f"expected 11 slides, got {n}")

    out = Presentation(BytesIO(data))
    w = out.slide_width / 914400
    h = out.slide_height / 914400
    print(f"output size_in={w:.2f}x{h:.2f}")
    if abs(w - 20.0) > 0.05 or abs(h - 11.25) > 0.05:
        raise AssertionError(f"expected 20x11.25, got {w}x{h}")

    checks = [
        (0, 8, "IANUSTEC"),
        (0, 9, "SOLUZIONI IT"),
        (0, 10, "2025"),
        (0, 11, "BY IANUSTEC TEAM"),
        (1, 3, "Indice"),
        (2, 11, "Chi Siamo"),
        (5, 8, "Visione e Valori"),
        (6, 9, "Il Problema"),
        (7, 3, "Il Nostro Team"),
        (8, 2, "I Nostri Servizi"),
        (10, 2, "Grazie!"),
    ]
    failed_text = []
    for sidx, sid, expect in checks:
        got = _shape_text(out.slides[sidx], sid, mod)
        ok = got is not None and expect in got
        print(f"  slide{sidx} id={sid} expect={expect!r} got={got!r} {'OK' if ok else 'FAIL'}")
        if not ok:
            failed_text.append((sidx, sid, expect, got))

    mstats = [_slide_stats(ref, i) for i in range(11)]
    ostata = [_slide_stats(data, i) for i in range(11)]
    reuse_slides = {
        i for i, s in enumerate(spec.get("slides") or []) if s.get("reuse") is not None
    }
    print("\n=== structure Marketing vs output ===")
    structure_fails: list[str] = []
    for i in range(11):
        a, b = mstats[i], ostata[i]
        print(
            f"  [{i}] mkt grp={a['grp']} cust={a['cust']} blip={a['blip']} svg={a['svg']}"
            f" embed={a['blip_embed']} spc={','.join(a['spc']) or '-'}"
            f"  | out grp={b['grp']} cust={b['cust']} blip={b['blip']} svg={b['svg']}"
            f" embed={b['blip_embed']} spc={','.join(b['spc']) or '-'}"
        )
        if b["dangling"]:
            structure_fails.append(f"slide {i}: dangling r:embed {b['dangling']}")
        if i not in reuse_slides:
            continue
        # A reuse slide is a graphic clone: groups, freeforms, blip fills and
        # letter-spacing must survive untouched (only the text changes).
        for key in ("grp", "cust", "blip", "svg", "blip_embed", "spc"):
            if a[key] != b[key]:
                structure_fails.append(
                    f"slide {i}: {key} marketing={a[key]} output={b[key]}"
                )

    cover_blips = ostata[0]["blip_embed"]
    mkt_cover_blips = mstats[0]["blip_embed"]
    print(
        f"\ncover blip embeds: marketing={mkt_cover_blips} output={cover_blips}"
    )

    # Slide 9: template table (id 2) refilled in place + title (id 9).
    tbl_shape = mod._shape_by_id(out.slides[9], 2)
    if tbl_shape is None or not tbl_shape.has_table:
        structure_fails.append("slide 9: template table id 2 not cloned")
    else:
        tbl = tbl_shape.table
        cells = [[c.text_frame.text.strip() for c in r.cells] for r in tbl.rows]
        spec_tbl = spec["slides"][9]["reuse"]["text"]["2"]
        want = [spec_tbl["headers"]] + spec_tbl["rows"]
        print(f"  slide9 table {len(cells)}x{len(cells[0]) if cells else 0}: {cells[0]} / {cells[1][:2]}")
        if cells != want:
            structure_fails.append(f"slide 9: table cells differ: {cells[:2]}")
        mkt_tbl = Presentation(BytesIO(ref)).slides[9].shapes
        mkt_tbl = [s for s in mkt_tbl if s.has_table][0]
        if abs(tbl_shape.width - mkt_tbl.width) > 2:
            structure_fails.append("slide 9: table width changed")
        # header cell formatting (bold, white, spc) must survive
        raw = zipfile.ZipFile(BytesIO(data)).read("ppt/slides/slide10.xml").decode()
        if 'b="true" sz="1766" spc="164"' not in raw and 'b="1" sz="1766" spc="164"' not in raw:
            structure_fails.append("slide 9: header run formatting lost")
        stale = [t for t in ("PHASE", "Phase 01", "Project Delivery Tracker") if t in raw]
        if stale:
            structure_fails.append(f"slide 9: template text left over: {stale}")
    title9 = _shape_text(out.slides[9], 9, mod)
    print(f"  slide9 id=9 got={title9!r}")
    if title9 != "Tracking Progetti":
        failed_text.append((9, 9, "Tracking Progetti", title9))

    path = asyncio.run(
        _save_via_cache(mod, data, spec.get("title") or "ianustec-reuse")
    )

    print("\n=== RESULT ===")
    if failed_text:
        print(f"TEXT FAILS: {len(failed_text)}")
        sys.exit(1)
    if structure_fails:
        print("STRUCTURE FAILS:")
        for line in structure_fails:
            print(f"  - {line}")
        sys.exit(1)
    print("PASS: 11 slides, 20x11.25, reuse texts, structure parity, table slide, cache save")
    print(f"OUTPUT: {path}")
    if cover_blips < mkt_cover_blips:
        print(
            "NOTE: cover lost blip fills vs Marketing "
            "(SVG/group clone fidelity — known gap)."
        )
    sys.exit(0)


if __name__ == "__main__":
    main()
