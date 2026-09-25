"""NF6: 15-slide deck in template mode opens without build errors.

    pip install -r examples/requirements-dev.txt
    python examples/template_mode_nf6_15_slides.py
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURE = BASE / "examples" / "fixtures" / "template_corners_one_slide.pptx"
OUT_DIR = BASE / "examples" / "output"
EXPECTED_SLIDES = 15


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)
    if not FIXTURE.is_file():
        print(f"Missing {FIXTURE}", file=sys.stderr)
        sys.exit(1)

    from pptx import Presentation

    source_bytes = FIXTURE.read_bytes()
    pack = mod._parse_reference_pptx(source_bytes)
    pack = mod._apply_template_edits(pack, None)

    slides = []
    for i in range(EXPECTED_SLIDES):
        if i % 2 == 0:
            slides.append(
                {
                    "layout": "title_bullets",
                    "title": f"Slide {i + 1}",
                    "bullets": [f"Bullet A on slide {i + 1}", "Bullet B"],
                }
            )
        else:
            slides.append(
                {
                    "layout": "title_body",
                    "title": f"Slide {i + 1}",
                    "body": f"Body paragraph for slide {i + 1}.\nSecond line.",
                }
            )

    spec = {
        "title": "NF6 template deck",
        "template_mapping": {"default": 0, "content": 0},
        "slides": slides,
    }

    data, n = mod.Tools()._build(
        spec,
        template_pack=pack,
        reference_bytes=source_bytes,
    )
    assert n == EXPECTED_SLIDES, n

    prs = Presentation(BytesIO(data))
    assert len(prs.slides) == EXPECTED_SLIDES

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / "template_mode_nf6_15_slides.pptx"
    out_path.write_bytes(data)
    prs2 = Presentation(str(out_path))
    assert len(prs2.slides) == EXPECTED_SLIDES

    print(f"PASS: NF6 {EXPECTED_SLIDES} slides, {len(data)} bytes -> {out_path}")


if __name__ == "__main__":
    main()
