"""Generate template-mode fixture .pptx files for local dev (Fase 0).

    pip install python-pptx pillow
    python examples/fixtures/build_fixtures.py

Idempotent: re-run overwrites the same paths under examples/fixtures/.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
OUT_DIR = Path(__file__).resolve().parent


def _blank_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def _add_blank_slide(prs: Presentation):
    layout = prs.slide_layouts[6]
    return prs.slides.add_slide(layout)


def _png_bytes(color: tuple[int, int, int], size: tuple[int, int] = (64, 64)) -> bytes:
    im = Image.new("RGB", size, color)
    buf = BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _corner_pictures(slide, size_in: float = 1.2) -> None:
    colors = [
        (200, 60, 60),
        (60, 120, 200),
        (60, 180, 100),
        (200, 160, 40),
    ]
    positions = [
        (0.15, 0.15),
        (SLIDE_W_IN - size_in - 0.15, 0.15),
        (0.15, SLIDE_H_IN - size_in - 0.15),
        (SLIDE_W_IN - size_in - 0.15, SLIDE_H_IN - size_in - 0.15),
    ]
    for (x, y), rgb in zip(positions, colors):
        slide.shapes.add_picture(
            BytesIO(_png_bytes(rgb)),
            Inches(x),
            Inches(y),
            width=Inches(size_in),
            height=Inches(size_in),
        )


def _white_safe_zone(slide) -> None:
    margin_x = 2.0
    margin_y = 1.8
    w = SLIDE_W_IN - 2 * margin_x
    h = SLIDE_H_IN - 2 * margin_y
    box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(margin_x),
        Inches(margin_y),
        Inches(w),
        Inches(h),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.line.fill.background()


def build_corners_one_slide(path: Path) -> None:
    prs = _blank_prs()
    slide = _add_blank_slide(prs)
    _white_safe_zone(slide)
    _corner_pictures(slide)
    prs.save(str(path))


def build_bg_fill_one_slide(path: Path) -> None:
    prs = _blank_prs()
    slide = _add_blank_slide(prs)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(230, 236, 245)
    slide.shapes.add_picture(
        BytesIO(_png_bytes((80, 80, 160), (128, 128))),
        Inches(SLIDE_W_IN - 2.5),
        Inches(0.4),
        width=Inches(2.0),
        height=Inches(2.0),
    )
    prs.save(str(path))


def build_two_slides(path: Path) -> None:
    prs = _blank_prs()
    cover = _add_blank_slide(prs)
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor(30, 40, 55)
    tb = cover.shapes.add_textbox(
        Inches(1.0), Inches(2.8), Inches(SLIDE_W_IN - 2.0), Inches(1.5)
    )
    tf = tb.text_frame
    tf.text = "Cover title (fixture)"
    p = tf.paragraphs[0]
    p.font.size = Pt(44)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    content = _add_blank_slide(prs)
    _white_safe_zone(content)
    _corner_pictures(content, size_in=0.9)
    body = content.shapes.add_textbox(
        Inches(2.2), Inches(2.0), Inches(SLIDE_W_IN - 4.4), Inches(3.5)
    )
    body.text_frame.text = "Content slide safe zone"
    prs.save(str(path))


def build_logo_confidential(path: Path) -> None:
    prs = _blank_prs()
    slide = _add_blank_slide(prs)
    slide.shapes.add_picture(
        BytesIO(_png_bytes((20, 20, 20), (96, 96))),
        Inches(0.5),
        Inches(0.35),
        width=Inches(0.9),
        height=Inches(0.9),
    )
    conf = slide.shapes.add_textbox(
        Inches(SLIDE_W_IN - 2.4), Inches(0.35), Inches(2.0), Inches(0.5)
    )
    conf.text_frame.text = "Confidential"
    conf.text_frame.paragraphs[0].font.size = Pt(14)
    conf.text_frame.paragraphs[0].font.color.rgb = RGBColor(120, 0, 0)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5),
        Inches(1.35),
        Inches(SLIDE_W_IN - 1.0),
        Inches(0.03),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(180, 180, 180)
    line.line.fill.background()
    prs.save(str(path))


def build_master_heavy_note(path: Path) -> None:
    """Minimal slide; v1 template mode does not clone slide masters."""
    prs = _blank_prs()
    slide = _add_blank_slide(prs)
    tb = slide.shapes.add_textbox(
        Inches(1.0), Inches(3.0), Inches(SLIDE_W_IN - 2.0), Inches(1.0)
    )
    tb.text_frame.text = "Master-heavy layouts: out of scope for template v1"
    prs.save(str(path))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    targets = {
        "template_corners_one_slide.pptx": build_corners_one_slide,
        "template_bg_fill_one_slide.pptx": build_bg_fill_one_slide,
        "template_two_slides.pptx": build_two_slides,
        "template_logo_confidential.pptx": build_logo_confidential,
        "template_master_heavy.pptx": build_master_heavy_note,
    }
    for name, fn in targets.items():
        out = OUT_DIR / name
        fn(out)
        print(f"Wrote {out}")


if __name__ == "__main__":
    main()
