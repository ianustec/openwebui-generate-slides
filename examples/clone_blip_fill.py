"""Clone an autoshape whose fill is an embedded image (a:blip).

Marketing-style freeforms store photos as blip fills, not as p:pic.
get_or_add_image_part returns (part, rId); treating that tuple as a Part
made save() raise AssertionError.

    pip install -r examples/requirements-dev.txt
    python examples/clone_blip_fill.py
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

BASE = Path(__file__).resolve().parent.parent

# 1x1 PNG
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
    "0000000c4944415408d763f8cfc00000030101005fe02b8a0000000049454e44ae426082"
)


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _reference_with_blip_fill() -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(
        BytesIO(_PNG), Inches(0), Inches(0), Inches(0.4), Inches(0.4)
    )
    blip = pic._element.find(".//" + qn("a:blip"))
    rid = blip.get(qn("r:embed"))
    pic._element.getparent().remove(pic._element)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2)
    )
    sp_pr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        node = sp_pr.find(qn(tag))
        if node is not None:
            sp_pr.remove(node)
    fill = etree.fromstring(
        f'<a:blipFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        f' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
    )
    sp_pr.append(fill)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> None:
    mod = _load_mod()
    raw = _reference_with_blip_fill()
    pack = mod._parse_reference_pptx(raw)
    fills = [
        d for st in pack.slides for d in st.decorations if d.kind != "picture"
    ]
    if not fills:
        raise SystemExit("FAIL: blip-fill autoshape was not collected as a decoration")
    spec = {
        "title": "Blip fill",
        "template_mapping": {"default": 0, "cover": 0},
        "template_edits": {"defaults": {"drop_text": True, "drop_offslide": True}},
        "slides": [{"layout": "cover", "title": "Clone", "subtitle": "blip"}],
    }
    data, n = mod.Tools()._build(spec, template_pack=pack, reference_bytes=raw)
    if n != 1:
        raise SystemExit(f"FAIL: expected 1 slide, got {n}")
    out = Presentation(BytesIO(data))
    blips = list(out.slides[0].shapes._spTree.iter(qn("a:blip")))
    if not blips:
        raise SystemExit("FAIL: cloned slide has no a:blip")
    part = out.slides[0].part
    for blip in blips:
        embed = blip.get(qn("r:embed"))
        if not embed:
            continue
        image = part.related_part(embed)
        if not getattr(image, "blob", None):
            raise SystemExit(f"FAIL: blip {embed} has no image blob")
    print("PASS: blip-fill autoshape cloned and package saved")


if __name__ == "__main__":
    sys.exit(main() or 0)
