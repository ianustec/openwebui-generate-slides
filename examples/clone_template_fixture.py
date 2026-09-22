"""Clone one template fixture slide to a new .pptx (Fase 3 dev check).

    pip install -r examples/requirements-dev.txt
    python examples/clone_template_fixture.py

Optional visual check: open examples/output/clone_corners.pptx in LibreOffice.
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"
OUT_DIR = BASE / "examples" / "output"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _picture_shapes(slide, mod) -> list:
    out = []
    for sh in slide.shapes:
        if mod._shape_kind(sh) == "picture":
            out.append(sh)
    return out


def _emu_tuple(shape) -> tuple[int, int, int, int]:
    return (int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    from pptx import Presentation

    corners = FIXTURES / "template_corners_one_slide.pptx"
    if not corners.is_file():
        print(f"Missing {corners}", file=sys.stderr)
        sys.exit(1)

    source_bytes = corners.read_bytes()
    source_prs = Presentation(BytesIO(source_bytes))
    pack = mod._parse_reference_pptx(source_bytes)
    pack = mod._apply_template_edits(pack, None)
    decs = pack.slides[0].decorations

    out_prs = Presentation()
    mod._clone_template_slide_to_prs(
        out_prs,
        source_prs,
        pack,
        source_slide_index=0,
        decorations=decs,
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / "clone_corners.pptx"
    out_prs.save(str(out_path))

    reread = Presentation(BytesIO(out_path.read_bytes()))
    assert len(reread.slides) == 1
    out_slide = reread.slides[0]
    pics = _picture_shapes(out_slide, mod)
    assert len(pics) >= 4, f"expected >=4 pictures, got {len(pics)}"

    source_slide = source_prs.slides[0]
    for dec in decs:
        if dec.kind != "picture":
            continue
        src = mod._shape_by_id(source_slide, dec.shape_id)
        if src is None:
            continue
        src_emu = _emu_tuple(src)
        if not any(_emu_tuple(p) == src_emu for p in pics):
            print(
                f"WARN: no matching picture bbox for source id={dec.shape_id} emu={src_emu}",
                file=sys.stderr,
            )

    print(f"OK: wrote {out_path} ({len(pics)} pictures)")

    pack_empty = mod._apply_template_edits(
        mod._parse_reference_pptx(source_bytes),
        {"slides": [{"index": 0, "keep_ids": []}]},
    )
    assert pack_empty.slides[0].decorations == [], "keep_ids [] should clear decorations"

    out2 = Presentation()
    mod._clone_template_slide_to_prs(
        out2,
        source_prs,
        pack_empty,
        source_slide_index=0,
        decorations=pack_empty.slides[0].decorations,
    )
    pics2 = _picture_shapes(out2.slides[0], mod)
    assert len(pics2) == 0, f"keep_ids [] should clone 0 pictures, got {len(pics2)}"
    print("OK: keep_ids [] → no cloned pictures")

    print("\nOK: clone fixture tests passed")
    sys.exit(0)


if __name__ == "__main__":
    main()
