"""Fase 3.4 R8: template_edits drop_text / drop_ids vs clone (offline).

    pip install -r examples/requirements-dev.txt
    python examples/inspect_template_text_contract.py
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"
LOGO = FIXTURES / "template_logo_confidential.pptx"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _confidential_hits(slide, mod) -> int:
    n = 0
    for shape in slide.shapes:
        txt = mod._shape_text(shape)
        if txt and "Confidential" in txt:
            n += 1
    return n


def _clone_with_edits(mod, source_bytes: bytes, edits) -> tuple:
    from pptx import Presentation

    source_prs = Presentation(BytesIO(source_bytes))
    pack = mod._apply_template_edits(mod._parse_reference_pptx(source_bytes), edits)
    out_prs = Presentation()
    mod._clone_template_slide_to_prs(
        out_prs,
        source_prs,
        pack,
        source_slide_index=0,
        decorations=pack.slides[0].decorations,
    )
    return out_prs, pack


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)
    if not LOGO.is_file():
        print(f"Missing {LOGO}", file=sys.stderr)
        sys.exit(1)

    data = LOGO.read_bytes()
    pack0 = mod._parse_reference_pptx(data)
    conf_id = next(
        (sh.id for sh in pack0.slides[0].shapes if sh.text and "Confidential" in sh.text),
        None,
    )
    if conf_id is None:
        print("FAIL: could not find Confidential shape id", file=sys.stderr)
        sys.exit(1)

    failed = False

    out_default, pack_d = _clone_with_edits(mod, data, None)
    if _confidential_hits(out_default.slides[0], mod) < 1:
        print("FAIL: default edits should clone Confidential text (drop_text false)", file=sys.stderr)
        failed = True
    else:
        print("OK default (None edits): Confidential text cloned")

    out_explicit, _ = _clone_with_edits(
        mod,
        data,
        {"defaults": {"drop_text": False, "drop_placeholders": False}},
    )
    if _confidential_hits(out_explicit.slides[0], mod) < 1:
        print("FAIL: explicit drop_text false should clone Confidential", file=sys.stderr)
        failed = True

    out_drop_text, _ = _clone_with_edits(mod, data, {"defaults": {"drop_text": True}})
    if _confidential_hits(out_drop_text.slides[0], mod) != 0:
        print("FAIL: drop_text true should omit template text", file=sys.stderr)
        failed = True
    else:
        print("OK drop_text true: no Confidential in clone")

    out_drop_id, _ = _clone_with_edits(
        mod,
        data,
        {"slides": [{"index": 0, "drop_ids": [conf_id]}]},
    )
    if _confidential_hits(out_drop_id.slides[0], mod) != 0:
        print(f"FAIL: drop_ids [{conf_id}] should remove Confidential", file=sys.stderr)
        failed = True
    else:
        print(f"OK drop_ids [{conf_id}]: Confidential removed")

    decs = pack_d.slides[0].decorations
    if not any(d.shape_id == conf_id for d in decs):
        print("FAIL: decorations should include Confidential shape when drop_text false", file=sys.stderr)
        failed = True

    if failed:
        sys.exit(1)
    print("\nOK: inspect_template_text_contract (R8 / F3 partial)")
    sys.exit(0)


if __name__ == "__main__":
    main()
