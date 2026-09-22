"""Parse template fixtures and print inventory (Fase 1 dev check).

    pip install -r examples/requirements-dev.txt
    python examples/parse_template_fixtures.py
"""
from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _print_pack(name: str, pack) -> None:
    print(f"\n=== {name} ===")
    print(
        f"  size {pack.slide_width_in:.3f}x{pack.slide_height_in:.3f} in, "
        f"slides={len(pack.slides)}, decorations_source={pack.decorations_source!r}"
    )
    print(f"  theme dk1={pack.theme.get('dk1')} accent1={pack.theme.get('accent1')}")
    for st in pack.slides:
        sz = st.safe_zone
        sz_s = (
            f"({sz.x:.2f},{sz.y:.2f},{sz.w:.2f},{sz.h:.2f})"
            if sz
            else "none"
        )
        print(
            f"  slide[{st.index}] shapes={st.shape_count} pics={st.picture_count} "
            f"decor={len(st.decorations)} safe={sz_s}"
        )
        if st.text_verbatim:
            print(f"    text_verbatim: {st.text_verbatim!r}")
        for sh in st.shapes[:12]:
            b = sh.bbox
            print(
                f"    id={sh.id} kind={sh.kind} z={sh.z_order} "
                f"bbox=({b.x:.2f},{b.y:.2f},{b.w:.2f},{b.h:.2f}) "
                f"text={sh.text!r}"
            )
        if len(st.shapes) > 12:
            print(f"    ... +{len(st.shapes) - 12} shapes")


def _assert_fixtures(mod, packs: dict) -> None:
    corners = packs.get("template_corners_one_slide.pptx")
    if corners:
        st0 = corners.slides[0]
        full = corners.slide_width_in * corners.slide_height_in
        sz = st0.safe_zone
        assert st0.picture_count >= 4, "corners: expected >=4 pictures"
        assert sz and sz.w * sz.h < full * 0.95, "corners: safe zone should be smaller than slide"

    logo = packs.get("template_logo_confidential.pptx")
    if logo:
        texts = logo.slides[0].text_verbatim
        assert any("Confidential" in t for t in texts), f"logo: need Confidential, got {texts}"
        assert logo.slides[0].picture_count >= 1, "logo: expected >=1 picture"

    two = packs.get("template_two_slides.pptx")
    if two:
        assert len(two.slides) == 2, "two_slides: expected 2 slides"

    # _is_text_shape smoke on logo fixture
    if logo and _HAS_PPTX(mod):
        from io import BytesIO
        from pptx import Presentation

        data = (FIXTURES / "template_logo_confidential.pptx").read_bytes()
        prs = Presentation(BytesIO(data))
        slide = prs.slides[0]
        kinds = []
        for shape in slide.shapes:
            kinds.append((shape.name, mod._is_text_shape(shape), mod._shape_kind(shape)))
        assert any(t for _, t, _ in kinds), "expected at least one text shape"
        assert any(k == "picture" for _, _, k in kinds), "expected picture shape"


def _HAS_PPTX(mod) -> bool:
    return getattr(mod, "_HAS_PPTX", False)


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    pptx_files = sorted(FIXTURES.glob("*.pptx"))
    if not pptx_files:
        print(f"No fixtures in {FIXTURES}", file=sys.stderr)
        sys.exit(1)

    packs = {}
    for path in pptx_files:
        try:
            pack = mod._parse_reference_pptx(path.read_bytes())
            packs[path.name] = pack
            _print_pack(path.name, pack)
        except Exception as exc:
            print(f"FAIL parse {path.name}: {exc}", file=sys.stderr)
            sys.exit(1)

    try:
        _assert_fixtures(mod, packs)
    except AssertionError as exc:
        print(f"ASSERT FAIL: {exc}", file=sys.stderr)
        sys.exit(1)

    print("\nOK: all fixture parses and assertions passed")
    sys.exit(0)


if __name__ == "__main__":
    main()
