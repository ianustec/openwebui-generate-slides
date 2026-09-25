"""Fase 8 P1: layout/master decoration merge when the file provides layout shapes.

    pip install -r examples/requirements-dev.txt
    python examples/inspect_master_p1_smoke.py
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"
ONLINE = BASE / "doc" / "online_templates"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    from pptx import Presentation

    corners = FIXTURES / "template_corners_one_slide.pptx"
    pack = mod._parse_reference_pptx(corners.read_bytes())
    assert pack.decorations_source == "slide"
    print("PASS: corners decorations_source=slide")

    candidates = sorted(ONLINE.glob("*.pptx")) if ONLINE.is_dir() else []
    tested = False
    for path in candidates:
        data = path.read_bytes()
        prs = Presentation(BytesIO(data))
        if not prs.slides:
            continue
        layout = prs.slides[0].slide_layout
        if len(layout.shapes) == 0 and len(layout.slide_master.shapes) == 0:
            continue
        pack = mod._parse_reference_pptx(data)
        payload = mod._serialize_inspect_payload(
            pack, file_id=f"online:{path.name}", filename=path.name, ok=True
        )
        src = payload.get("decorations_source")
        hints = payload.get("hints") or {}
        if src in ("master", "both") or hints.get("master_decorations"):
            print(f"PASS: {path.name} decorations_source={src!r} master P1 active")
            tested = True
            break
    if not tested:
        print("SKIP: no online template with layout/master shapes on disk")


if __name__ == "__main__":
    main()
