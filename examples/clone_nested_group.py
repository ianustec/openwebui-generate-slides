"""Fase 8: flattened decorations align with inspect ids + recursive shape lookup.

    pip install -r examples/requirements-dev.txt
    python examples/clone_nested_group.py
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    from pptx import Presentation

    for name in ("template_nested_group.pptx", "template_corners_one_slide.pptx"):
        path = FIXTURES / name
        if not path.is_file():
            print(f"SKIP missing {path}")
            continue
        data = path.read_bytes()
        pack = mod._parse_reference_pptx(data)
        prs = Presentation(BytesIO(data))
        slide = prs.slides[0]
        st = pack.slides[0]
        dec_ids = {d.shape_id for d in st.decorations}
        for sid in dec_ids:
            if mod._shape_by_id_recursive(slide, sid) is None:
                print(f"FAIL: shape id {sid} not found recursively in {name}", file=sys.stderr)
                sys.exit(1)
        pic_ids = {s.id for s in st.shapes if s.kind == "picture"}
        if not pic_ids.issubset(dec_ids):
            print(
                f"FAIL: picture ids {pic_ids - dec_ids} missing from decorations in {name}",
                file=sys.stderr,
            )
            sys.exit(1)
        print(f"PASS: {name} decorations={len(dec_ids)} recursive lookup OK")

    corners = FIXTURES / "template_corners_one_slide.pptx"
    source_bytes = corners.read_bytes()
    pack = mod._parse_reference_pptx(source_bytes)
    spec = {
        "title": "Nested clone smoke",
        "template_mapping": {"default": 0, "content": 0},
        "slides": [
            {
                "layout": "title_bullets",
                "title": "Test",
                "bullets": ["A", "B"],
            }
        ],
    }
    out, n = mod.Tools()._build(
        spec,
        template_pack=pack,
        reference_bytes=source_bytes,
    )
    prs = Presentation(BytesIO(out))
    assert n == 1 and len(prs.slides) == 1
    print("PASS: template _build clone after flatten decorations")


if __name__ == "__main__":
    main()
