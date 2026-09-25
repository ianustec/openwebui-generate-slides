"""Per-slide reuse: clone template slide, replace text by shape id, no layout overlay.

    pip install -r examples/requirements-dev.txt
    python examples/reuse_slide_text.py
"""
from __future__ import annotations

import importlib.util
import json
import sys
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"
MARKETING = BASE / "doc" / "online_templates" / "Marketing.pptx"
OUT_DIR = BASE / "examples" / "output"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _fixture_with_textbox() -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tb.text_frame.paragraphs[0].text = "ORIGINAL_TITLE"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _shape_text(slide, shape_id: int) -> str | None:
    sh = mod._shape_by_id(slide, shape_id)
    if sh is None or not sh.has_text_frame:
        return None
    return (sh.text or "").strip()


def _table_cells(slide, shape_id: int) -> list[list[str]]:
    sh = mod._shape_by_id(slide, shape_id)
    assert sh is not None and sh.has_table, f"table id {shape_id} not cloned"
    return [[c.text_frame.text.strip() for c in r.cells] for r in sh.table.rows]


def _table_tests(mod, mpack, mraw: bytes) -> None:
    """Marketing slide 9: table id 2 (6x5), title id 9."""
    st = mpack.slides[9]
    rec = next(r for r in st.shapes if r.id == 2)
    assert rec.kind == "table" and rec.cloneable, rec
    assert rec.table and rec.table["rows"] == 6 and rec.table["cols"] == 5
    assert rec.table["cells"][0][0] == "PHASE"
    print("OK: inspect exposes table grid (6x5, cloneable)")

    src_tbl = [s for s in Presentation(BytesIO(mraw)).slides[9].shapes if s.has_table][0]

    # 1) explicit payload, fewer rows and fewer columns than the template
    spec = {
        "title": "t",
        "slides": [
            {
                "reuse": {
                    "slide": 9,
                    "text": {
                        "9": "Tracking",
                        "2": {
                            "headers": ["A", "B", "C"],
                            "rows": [["1", "2", "3"], ["4", "5", "6"]],
                        },
                    },
                }
            }
        ],
    }
    data, _ = mod.Tools()._build(spec, template_pack=mpack, reference_bytes=mraw)
    out = Presentation(BytesIO(data))
    cells = _table_cells(out.slides[0], 2)
    assert cells == [["A", "B", "C"], ["1", "2", "3"], ["4", "5", "6"]], cells
    out_tbl = mod._shape_by_id(out.slides[0], 2)
    assert abs(out_tbl.width - src_tbl.width) <= 2, "table width not preserved"
    assert out_tbl.height < src_tbl.height, "frame height should shrink with rows"
    assert _shape_text(out.slides[0], 9) == "Tracking"
    print("OK: table 3x3 refilled (cols/rows removed, width kept, height shrunk)")

    # 2) more rows and more columns than the template
    spec["slides"][0]["reuse"]["text"]["2"] = {
        "headers": [f"H{i}" for i in range(7)],
        "rows": [[f"r{r}c{c}" for c in range(7)] for r in range(9)],
    }
    data, _ = mod.Tools()._build(spec, template_pack=mpack, reference_bytes=mraw)
    out = Presentation(BytesIO(data))
    cells = _table_cells(out.slides[0], 2)
    assert len(cells) == 10 and all(len(r) == 7 for r in cells), (len(cells), [len(r) for r in cells])
    assert cells[0][6] == "H6" and cells[9][6] == "r8c6"
    out_tbl = mod._shape_by_id(out.slides[0], 2)
    assert abs(out_tbl.width - src_tbl.width) <= 2
    print("OK: table 10x7 refilled (cols/rows added, width kept)")

    # 3) slide-level headers/rows routed to the single template table
    spec_fallback = {
        "title": "t",
        "slides": [
            {
                "layout": "table",
                "title": "ignored",
                "headers": ["X", "Y"],
                "rows": [["1", "2"]],
                "reuse": {"slide": 9, "text": {"9": "Fallback"}},
            }
        ],
    }
    data, _ = mod.Tools()._build(spec_fallback, template_pack=mpack, reference_bytes=mraw)
    out = Presentation(BytesIO(data))
    assert _table_cells(out.slides[0], 2) == [["X", "Y"], ["1", "2"]]
    print("OK: slide-level headers/rows routed to template table")

    # 4) string payload on a table id is rejected with a clear message
    bad = {"title": "t", "slides": [{"reuse": {"slide": 9, "text": {"2": "oops"}}}]}
    try:
        mod.Tools()._build(bad, template_pack=mpack, reference_bytes=mraw)
        raise AssertionError("expected ValueError for string on table id")
    except ValueError as exc:
        assert "headers" in str(exc), str(exc)
    print("OK: string payload on table id rejected")

    # 5) explicit per-layout template_mapping key wins over role fallback
    assert mod._pick_template_slide(mpack, "table", {"content": 1, "table": 9}) == 9
    assert mod._pick_template_slide(mpack, "team", {"content": 1, "team": 7}) == 7
    assert mod._pick_template_slide(mpack, "title_body", {"content": 1}) == 1
    print("OK: template_mapping explicit layout keys honoured")


def main() -> None:
    global mod
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    real_id = "30dbda41-a9a7-4e60-a8be-00a9d538d818"
    assert not mod._is_files_api_id("20946")
    assert not mod._is_files_api_id("Indian Doctors Day.pptx")
    assert mod._is_files_api_id(real_id)
    invented = mod._invented_reference_id_message(
        "20946",
        [{"files": [{"id": real_id, "name": "Indian Doctors Day.pptx"}]}],
    )
    assert "20946" in invented and real_id in invented
    assert real_id not in mod._invented_reference_id_message("20946", [])
    wrong = mod._wrong_reference_id_message(
        "089603a1-9665-4443-ac56-b1500479f9ba", real_id
    )
    assert real_id in wrong and "089603a1" in wrong and "classic" in wrong.lower()
    compact = mod._compact_inspect_payload(
        {
            "ok": True,
            "file_id": real_id,
            "filename": "t.pptx",
            "slide_width_in": 10,
            "slide_height_in": 5.625,
            "slide_count": 1,
            "slides": [
                {
                    "index": 0,
                    "shape_count": 3,
                    "shapes": [
                        {"id": 1, "kind": "textbox", "has_text": True, "text": "TITLE"},
                        {"id": 2, "kind": "autoshape", "has_text": False, "text": None},
                        {
                            "id": 3,
                            "kind": "table",
                            "table": {"rows": 2, "cols": 2, "cells": [["A", "B"]]},
                        },
                    ],
                }
            ],
        }
    )
    assert compact["slides"][0]["text"] == [{"id": 1, "text": "TITLE"}]
    assert compact["slides"][0]["tables"][0]["id"] == 3
    assert compact["slides"][0]["decoration_count"] == 1
    assert "bbox" not in json.dumps(compact)
    assert real_id in compact["next"]
    print("OK: invented reference_file_id rejected with the attached UUID")

    raw = _fixture_with_textbox()
    pack = mod._parse_reference_pptx(raw)
    rec = pack.slides[0].shapes[0]
    sid = rec.id
    spec = {
        "title": "Reuse test",
        "slides": [
            {
                "layout": "cover",
                "title": "SHOULD_NOT_APPEAR",
                "reuse": {
                    "slide": 0,
                    "keep_ids": [sid],
                    "text": {str(sid): "REPLACED_TITLE"},
                },
            }
        ],
    }
    data, n = mod.Tools()._build(spec, template_pack=pack, reference_bytes=raw)
    assert n == 1
    out = Presentation(BytesIO(data))
    assert len(out.slides) == 1
    txt = _shape_text(out.slides[0], sid)
    assert txt == "REPLACED_TITLE", f"expected REPLACED_TITLE, got {txt!r}"
    all_text = " ".join(
        p.text for sh in out.slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs
    )
    assert "SHOULD_NOT_APPEAR" not in all_text
    print("OK: reuse replaces text in place, no cover overlay")

    try:
        mod._parse_reuse_spec(["0:1"], pack)
        raise AssertionError("expected ValueError for list reuse")
    except ValueError as exc:
        assert "object" in str(exc).lower() or "list" in str(exc).lower()
    print("OK: invalid reuse list rejected")

    bad_spec = {"title": "x", "slides": [{"reuse": {"slide": 99}}]}
    try:
        mod.Tools()._build(bad_spec, template_pack=pack, reference_bytes=raw)
        raise AssertionError("expected out of range error")
    except ValueError as exc:
        assert "out of range" in str(exc)
    print("OK: reuse.slide out of range fails")

    mixed = {
        "title": "mixed",
        "template_mapping": {"default": 0, "cover": 0},
        "slides": [
            {
                "reuse": {
                    "slide": 0,
                    "keep_ids": [sid],
                    "text": {str(sid): "REUSE_ONLY"},
                },
            },
            {"layout": "title_body", "title": "Classic overlay", "body": "body"},
        ],
    }
    data2, n2 = mod.Tools()._build(mixed, template_pack=pack, reference_bytes=raw)
    assert n2 == 2
    out2 = Presentation(BytesIO(data2))
    assert len(out2.slides) == 2
    print("OK: mixed deck reuse + classic mapping slide")

    if MARKETING.is_file():
        mraw = MARKETING.read_bytes()
        mpack = mod._parse_reference_pptx(mraw)
        mspec = {
            "title": "Marketing cover reuse",
            "slides": [
                {
                    "reuse": {
                        "slide": 0,
                        "drop_ids": [10],
                        "text": {
                            "8": "IANUSTEC",
                            "9": "PROPOSAL",
                            "11": "BY TEST",
                        },
                    },
                }
            ],
        }
        mdata, _ = mod.Tools()._build(
            mspec, template_pack=mpack, reference_bytes=mraw
        )
        mout = Presentation(BytesIO(mdata))
        assert round(mout.slide_width / 914400, 2) == 20.0
        assert round(mout.slide_height / 914400, 2) == 11.25
        t8 = _shape_text(mout.slides[0], 8)
        assert t8 and "IANUSTEC" in t8
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        (OUT_DIR / "reuse_marketing_cover_smoke.pptx").write_bytes(mdata)
        print("OK: Marketing cover reuse smoke (20x11.25, text id 8)")

        _table_tests(mod, mpack, mraw)
    else:
        print("SKIP: Marketing.pptx not found for smoke")

    print("\nOK: reuse_slide_text tests passed")
    sys.exit(0)


if __name__ == "__main__":
    main()
