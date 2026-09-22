"""Generate inspect_slides JSON for doc/online_templates/*.pptx and validate vs source.

    pip install -r examples/requirements-dev.txt
    python examples/inspect_online_templates.py

Writes JSON under doc/online_templates/inspect_output/ and _summary.json.
"""
from __future__ import annotations

import importlib.util
import json
import sys
import time
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
TEMPLATES = BASE / "doc" / "online_templates"
OUT = TEMPLATES / "inspect_output"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _pptx_all_text(prs) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[str] = []

    def walk(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for ch in shape.shapes:
                walk(ch)
            return
        if getattr(shape, "has_text_frame", False) and shape.text_frame:
            t = (shape.text_frame.text or "").strip()
            if t:
                out.append(t)

    for slide in prs.slides:
        for shape in slide.shapes:
            walk(shape)
    return out


def _json_all_text(payload: dict) -> list[str]:
    out: list[str] = []
    for slide in payload.get("slides") or []:
        out.extend(slide.get("text_verbatim") or [])
        for sh in slide.get("shapes") or []:
            t = (sh.get("text") or "").strip()
            if t:
                out.append(t)
    return out


def validate_structural(payload: dict, prs, mod) -> list[str]:
    errs: list[str] = []
    if payload.get("slide_count") != len(prs.slides):
        errs.append("slide_count mismatch")
    exp_w, exp_h = mod._emu_in(prs.slide_width), mod._emu_in(prs.slide_height)
    if abs(payload["slide_width_in"] - exp_w) > 0.02:
        errs.append("width mismatch")
    if abs(payload["slide_height_in"] - exp_h) > 0.02:
        errs.append("height mismatch")
    if payload.get("filename") and not payload["filename"].endswith(".pptx"):
        errs.append("filename should end with .pptx")
    for slide in payload.get("slides") or []:
        idx = slide["index"]
        shapes = slide.get("shapes") or []
        if slide["shape_count"] != len(shapes):
            errs.append(f"slide {idx} shape_count")
        pc = sum(1 for s in shapes if s.get("kind") == "picture")
        if slide["picture_count"] != pc:
            errs.append(f"slide {idx} picture_count")
        sz = slide.get("safe_zone")
        if sz and not sz.get("method"):
            errs.append(f"slide {idx} safe_zone.method")
    return errs


def _norm_text(s: str) -> str:
    return " ".join((s or "").split()).casefold()


def validate_semantic(payload: dict, prs) -> list[str]:
    errs: list[str] = []
    pptx_text = _pptx_all_text(prs)
    json_text = _json_all_text(payload)
    if not pptx_text and not json_text:
        return errs
    json_blob = _norm_text("\n".join(json_text))
    missing: list[str] = []
    for block in pptx_text:
        nb = _norm_text(block)
        if len(nb) < 12:
            continue
        if nb in json_blob:
            continue
        # Long Slidesgo boilerplate: require most lines to appear individually.
        lines = [_norm_text(ln) for ln in block.splitlines() if _norm_text(ln)]
        if lines:
            hit = sum(1 for ln in lines if ln in json_blob)
            if hit >= max(1, int(0.85 * len(lines))):
                continue
        missing.append(block[:120])
    if missing:
        errs.append(
            f"text blocks missing from JSON ({len(missing)}/{len(pptx_text)}), "
            f"sample: {missing[:3]!r}"
        )
    return errs


def main() -> None:
    from pptx import Presentation

    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    if not TEMPLATES.is_dir():
        print(f"Missing {TEMPLATES}", file=sys.stderr)
        sys.exit(1)

    OUT.mkdir(parents=True, exist_ok=True)
    files = sorted(TEMPLATES.glob("*.pptx"))
    if not files:
        print("No .pptx in online_templates", file=sys.stderr)
        sys.exit(1)

    summary: list[dict] = []
    for path in files:
        row: dict = {"file": path.name}
        t0 = time.time()
        try:
            data = path.read_bytes()
            prs = Presentation(BytesIO(data))
            pack = mod._parse_reference_pptx(data)
            payload = mod._serialize_inspect_payload(
                pack,
                file_id=f"local:{path.stem}",
                filename=path.name,
                ok=True,
            )
            struct_errs = validate_structural(payload, prs, mod)
            sem_errs = validate_semantic(payload, prs)
            errs = struct_errs + sem_errs
            out = OUT / f"{path.stem}.inspect.json"
            out.write_text(
                json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            row.update(
                {
                    "slide_count": payload["slide_count"],
                    "dimensions_in": [
                        payload["slide_width_in"],
                        payload["slide_height_in"],
                    ],
                    "shapes_total": sum(s["shape_count"] for s in payload["slides"]),
                    "pictures_total": sum(
                        s["picture_count"] for s in payload["slides"]
                    ),
                    "theme": payload.get("theme"),
                    "valid": len(errs) == 0,
                    "validation_errors": errs,
                    "json_path": str(out.relative_to(BASE)),
                    "elapsed_sec": round(time.time() - t0, 2),
                }
            )
            if payload["slides"]:
                s0 = payload["slides"][0]
                row["slide0"] = {
                    "shapes": s0["shape_count"],
                    "pictures": s0["picture_count"],
                    "text_head": (s0.get("text_verbatim") or [])[:5],
                    "kinds": sorted({sh["kind"] for sh in s0["shapes"]}),
                }
        except Exception as exc:
            row["valid"] = False
            row["error"] = str(exc)
            row["elapsed_sec"] = round(time.time() - t0, 2)
        summary.append(row)
        status = "OK" if row.get("valid") else "FAIL"
        print(f"{status} {path.name} ({row.get('elapsed_sec')}s)")

    (OUT / "_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    failed = sum(1 for r in summary if not r.get("valid"))
    print(f"\n{len(summary) - failed}/{len(summary)} passed → {OUT / '_summary.json'}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
