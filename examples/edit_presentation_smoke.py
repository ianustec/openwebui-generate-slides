"""Smoke tests for post-generate presentation edit (offline, no Open WebUI).

    pip install -r examples/requirements-dev.txt
    python examples/edit_presentation_smoke.py
"""
from __future__ import annotations

import asyncio
import importlib.util
import json
import sys
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent.parent
MARKETING = BASE / "doc" / "online_templates" / "Marketing.pptx"
VALID_UUID = "8f3c2a1b-1234-5678-9abc-def012345678"


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(mod)
    return mod


class _ValvesStub:
    presentation_edit_min_font_pt = 10.0
    presentation_edit_shrink_step_pt = 2.0


def _edit_json_slide(internal_index: int = 0) -> int:
    """`operations[].slide` as in edit_presentation JSON (see EDIT_SLIDE_INDEX_ORIGIN)."""
    origin = getattr(mod, "EDIT_SLIDE_INDEX_ORIGIN", 0)
    return internal_index + (1 if origin == 1 else 0)


def _build_overflow_fixture() -> tuple[bytes, int]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2.5), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = (
        "OVERFLOW TITLE TEXT THAT DOES NOT FIT IN A SMALL BOX "
        "WITHOUT REDUCING FONT SIZE FOR PRESENTATION EDIT SMOKE"
    )
    p.runs[0].font.size = Pt(36)
    p.runs[0].font.bold = True
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue(), tb.shape_id


def _build_spc_fixture() -> tuple[bytes, int, str]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    p = tb.text_frame.paragraphs[0]
    p.text = "ORIGINAL"
    run = p.runs[0]
    run.font.size = Pt(24)
    r_pr = run._r.get_or_add_rPr()
    r_pr.set("spc", "800")
    r_pr.set("b", "1")
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue(), tb.shape_id, "800"


def _slide_spc_values(slide, shape_id: int) -> list[str]:
    sh = mod._shape_by_id_recursive(slide, shape_id)
    assert sh is not None and sh.has_text_frame
    tx = sh.text_frame._txBody
    out: list[str] = []
    for p in tx.findall(mod.qn("a:p")):
        for r in p.findall(mod.qn("a:r")):
            r_pr = r.find(mod.qn("a:rPr"))
            if r_pr is not None and r_pr.get("spc") is not None:
                out.append(r_pr.get("spc"))
    return out


def _slide_bold(slide, shape_id: int) -> list[str | None]:
    sh = mod._shape_by_id_recursive(slide, shape_id)
    assert sh is not None
    tx = sh.text_frame._txBody
    vals: list[str | None] = []
    for p in tx.findall(mod.qn("a:p")):
        for r in p.findall(mod.qn("a:r")):
            r_pr = r.find(mod.qn("a:rPr"))
            vals.append(r_pr.get("b") if r_pr is not None else None)
    return vals


def _test_shrink_pe1(mod) -> None:
    data, sid = _build_overflow_fixture()
    prs = Presentation(BytesIO(data))
    shape = mod._shape_by_id_recursive(prs.slides[0], sid)
    assert shape is not None
    before = mod._edit_shape_font_pt_max(shape)
    assert before >= 30
    mod._apply_edit_operations(
        prs,
        [
            {
                "op": "shrink_font",
                "slide": 0,
                "shape_id": sid,
                "min_pt": 10,
                "step_pt": 2,
                "max_iterations": 20,
            }
        ],
        _ValvesStub(),
    )
    after = mod._edit_shape_font_pt_max(shape)
    assert after < before, f"expected sz decrease: {before} -> {after}"
    print(f"OK: shrink_font reduced font {before}pt -> {after}pt (PE1)")


def _test_set_text_pe2(mod) -> None:
    data, sid, spc = _build_spc_fixture()
    prs = Presentation(BytesIO(data))
    mod._apply_edit_operations(
        prs,
        [
            {
                "op": "set_text",
                "slide": 0,
                "shape_id": sid,
                "text": "UPDATED\nLINE2",
            }
        ],
        _ValvesStub(),
    )
    sh = mod._shape_by_id_recursive(prs.slides[0], sid)
    assert sh is not None
    assert "UPDATED" in (sh.text or "")
    assert _slide_spc_values(prs.slides[0], sid) == [spc, spc]
    assert all(v == "1" for v in _slide_bold(prs.slides[0], sid))
    print("OK: set_text preserves spc and bold (PE2)")


def _test_parse_and_apply_errors(mod) -> None:
    v = _ValvesStub()
    try:
        mod._parse_edit_spec({"operations": []}, v)
        raise AssertionError("expected empty operations error")
    except ValueError:
        pass
    try:
        mod._parse_edit_spec(
            {
                "presentation_file_id": "not-a-uuid",
                "operations": [
                    {
                        "op": "set_text",
                        "slide": _edit_json_slide(0),
                        "shape_id": 1,
                        "text": "x",
                    }
                ],
            },
            v,
        )
        raise AssertionError("expected bad uuid error")
    except ValueError as exc:
        assert "Files API" in str(exc) or "UUID" in str(exc)
    try:
        mod._parse_edit_spec({}, v)
        raise AssertionError("expected missing id")
    except ValueError:
        pass

    data, sid = _build_overflow_fixture()
    prs = Presentation(BytesIO(data))
    try:
        mod._apply_edit_operations(
            prs,
            [{"op": "set_text", "slide": 99, "shape_id": sid, "text": "x"}],
            v,
        )
        raise AssertionError("expected slide oob")
    except ValueError as exc:
        msg = str(exc)
        assert "operations[0]" in msg and "out of range" in msg

    try:
        mod._apply_edit_operations(
            prs,
            [{"op": "set_text", "slide": 0, "shape_id": 999999, "text": "x"}],
            v,
        )
        raise AssertionError("expected missing shape")
    except ValueError as exc:
        msg = str(exc)
        assert "operations[0]" in msg and "not found" in msg

    print("OK: parse UUID + apply slide/shape errors (PE5, PE6, PE8)")


async def _test_edit_tool_guards(mod) -> None:
    tools = mod.Tools()
    tools.valves.presentation_edit_enabled = False
    reply = await tools.edit_presentation("{}")
    low = reply.lower()
    assert "couldn't edit" in low
    assert "presentation_edit_enabled" in low

    tools.valves.presentation_edit_enabled = True
    bad = await tools.edit_presentation(
        json.dumps(
            {
                "presentation_file_id": VALID_UUID,
                "operations": [
                    {
                        "op": "set_text",
                        "slide": _edit_json_slide(0),
                        "shape_id": 1,
                        "text": "x",
                    }
                ],
            }
        )
    )
    assert "couldn't edit" in bad.lower()
    assert "user" in bad.lower() or "load" in bad.lower() or "file" in bad.lower()

    print("OK: edit_presentation valve off + load fail without OWUI (PE4)")


def _test_marketing_reuse_post_edit(mod) -> None:
    if not MARKETING.is_file():
        print(f"SKIP Marketing reuse edit (missing {MARKETING})")
        return
    mraw = MARKETING.read_bytes()
    mpack = mod._parse_reference_pptx(mraw)
    long_title = (
        "EXTENDED MARKETING TITLE FOR SHRINK TEST "
        "THAT SHOULD NOT FIT THE TEMPLATE PLACEHOLDER"
    )
    spec = {
        "title": "edit-smoke-marketing",
        "slides": [
            {
                "reuse": {
                    "slide": 0,
                    "text": {"8": long_title, "9": "SUB", "11": "FOOTER"},
                },
            }
        ],
    }
    data, n = mod.Tools()._build(
        spec, template_pack=mpack, reference_bytes=mraw
    )
    assert n == 1
    prs = Presentation(BytesIO(data))
    sid = 8
    shape = mod._shape_by_id_recursive(prs.slides[0], sid)
    assert shape is not None and shape.has_text_frame
    before = mod._edit_shape_font_pt_max(shape)
    mod._apply_edit_operations(
        prs,
        [
            {
                "op": "shrink_font",
                "slide": 0,
                "shape_id": sid,
                "min_pt": 8,
                "step_pt": 2,
                "max_iterations": 25,
            }
        ],
        _ValvesStub(),
    )
    after = mod._edit_shape_font_pt_max(shape)
    assert after <= before
    if after < before:
        print(f"OK: Marketing reuse post-edit shrink {before}pt -> {after}pt")
    else:
        # Heuristic may already fit at template size; still exercised path.
        assert long_title in (shape.text or "")
        print("OK: Marketing reuse post-edit (shrink no-op at min/template size)")


def _build_split_fixture() -> tuple[bytes, int, int]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1.5))
    b = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(4), Inches(1.5))
    a.text_frame.paragraphs[0].text = "TITLE LINE"
    a.text_frame.add_paragraph().text = "Body paragraph here"
    b.text_frame.paragraphs[0].text = ""
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue(), a.shape_id, b.shape_id


def _build_table_fixture() -> tuple[bytes, int]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1.5))
    tbl.table.cell(0, 0).text = "H1"
    tbl.table.cell(1, 1).text = "OLD"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue(), tbl.shape_id


def _test_v2_ops(mod) -> None:
    v = _ValvesStub()
    data, sid = _build_overflow_fixture()
    prs = Presentation(BytesIO(data))
    shape = mod._shape_by_id_recursive(prs.slides[0], sid)
    before = mod._edit_shape_font_pt_max(shape)
    mod._apply_edit_operations(
        prs,
        [
            {
                "op": "replace_text_and_fit",
                "slide": 0,
                "shape_id": sid,
                "text": "SHORTER TITLE TEXT",
                "min_pt": 10,
                "step_pt": 2,
                "max_iterations": 20,
            }
        ],
        v,
    )
    assert "SHORTER" in (shape.text or "")
    assert mod._edit_shape_font_pt_max(shape) <= before
    print("OK: replace_text_and_fit (v2)")

    data2, sid2 = _build_overflow_fixture()
    prs2 = Presentation(BytesIO(data2))
    mod._apply_edit_operations(
        prs2,
        [{"op": "enable_autofit", "slide": 0, "shape_id": sid2}],
        v,
    )
    sh2 = mod._shape_by_id_recursive(prs2.slides[0], sid2)
    body = sh2.text_frame._txBody.find(mod.qn("a:bodyPr"))
    assert body is not None
    assert body.find(mod.qn("a:normAutofit")) is not None
    assert body.find(mod.qn("a:noAutofit")) is None
    print("OK: enable_autofit normAutofit (v2)")

    data3, id_a, id_b = _build_split_fixture()
    prs3 = Presentation(BytesIO(data3))
    mod._apply_edit_operations(
        prs3,
        [
            {
                "op": "split_text",
                "slide": 0,
                "from_shape_id": id_a,
                "to_shape_id": id_b,
                "mode": "first_line",
            }
        ],
        v,
    )
    sa = mod._shape_by_id_recursive(prs3.slides[0], id_a)
    sb = mod._shape_by_id_recursive(prs3.slides[0], id_b)
    assert (sa.text or "").strip() == "TITLE LINE"
    assert "Body paragraph" in (sb.text or "")
    print("OK: split_text first_line (v2)")

    data4, sid4 = _build_overflow_fixture()
    prs4 = Presentation(BytesIO(data4))
    sh4 = mod._shape_by_id_recursive(prs4.slides[0], sid4)
    h0 = int(sh4.height)
    resize_op = mod._normalize_edit_op(
        {
            "op": "resize_shape",
            "slide": _edit_json_slide(0),
            "shape_id": sid4,
            "delta_height_in": 0.2,
        },
        0,
        v,
    )
    mod._apply_edit_operations(prs4, [resize_op], v)
    assert int(sh4.height) > h0
    print("OK: resize_shape delta_height_in (v2)")

    data5, tid = _build_table_fixture()
    prs5 = Presentation(BytesIO(data5))
    mod._apply_edit_operations(
        prs5,
        [
            {
                "op": "set_table_cell",
                "slide": 0,
                "shape_id": tid,
                "row": 1,
                "col": 1,
                "text": "NEW",
            }
        ],
        v,
    )
    sh5 = mod._shape_by_id_recursive(prs5.slides[0], tid)
    assert sh5.table.cell(1, 1).text.strip() == "NEW"
    print("OK: set_table_cell (v2)")


def _test_marketing_autofit_and_inspect(mod) -> None:
    if not MARKETING.is_file():
        print(f"SKIP Marketing v2 autofit/inspect (missing {MARKETING})")
        return
    mraw = MARKETING.read_bytes()
    mpack = mod._parse_reference_pptx(mraw)
    spec = {
        "title": "inspect-smoke",
        "slides": [{"reuse": {"slide": 0, "text": {"8": "A", "9": "B"}}}],
    }
    deck_bytes, _n = mod.Tools()._build(
        spec, template_pack=mpack, reference_bytes=mraw
    )
    pack_out = mod._parse_reference_pptx(deck_bytes)
    payload = mod._serialize_inspect_payload(
        pack_out, file_id=VALID_UUID, filename="out.pptx", ok=True
    )
    compact = mod._compact_inspect_payload(payload)
    texts = compact["slides"][0].get("text") or []
    assert len(texts) >= 2, "compact inspect needs >=2 text ids for split workflow"
    for t in texts:
        assert t.get("id") is not None and "text" in t
    print(f"OK: compact inspect 10A ({len(texts)} text shapes on Marketing reuse)")

    prs = Presentation(BytesIO(deck_bytes))
    mod._apply_edit_operations(
        prs,
        [{"op": "enable_autofit", "slide": 0, "shape_id": 8}],
        _ValvesStub(),
    )
    sh = mod._shape_by_id_recursive(prs.slides[0], 8)
    body = sh.text_frame._txBody.find(mod.qn("a:bodyPr"))
    assert body is not None and body.find(mod.qn("a:normAutofit")) is not None
    print("OK: Marketing reuse enable_autofit (v2)")


def _build_endpara_sz_fixture() -> tuple[bytes, int]:
    """Text shape with font size only on a:endParaRPr (no run sz)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(2), Inches(1), Inches(5), Inches(1.2))
    tb.text_frame.paragraphs[0].text = "TITLE FROM ENDPARA"
    tx = tb.text_frame._txBody
    p = tx.findall(mod.qn("a:p"))[0]
    for r in list(p.findall(mod.qn("a:r"))):
        r_pr = r.find(mod.qn("a:rPr"))
        if r_pr is not None and r_pr.get("sz") is not None:
            r_pr.attrib.pop("sz", None)
    end = p.find(mod.qn("a:endParaRPr"))
    if end is None:
        from lxml import etree

        end = etree.SubElement(p, mod.qn("a:endParaRPr"))
    end.set("sz", "3600")
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue(), tb.shape_id


def _test_endpara_effective_and_to_min(mod) -> None:
    data, sid = _build_endpara_sz_fixture()
    prs = Presentation(BytesIO(data))
    shape = mod._shape_by_id_recursive(prs.slides[0], sid)
    assert shape is not None
    eff = mod._edit_shape_font_pt_effective(shape)
    assert eff >= 35.0, f"expected endParaRPr 36pt effective, got {eff}"
    mutated = mod._apply_edit_operations(
        prs,
        [
            {
                "op": "shrink_font",
                "slide": 0,
                "shape_id": sid,
                "shrink_mode": "to_min",
                "min_pt": 20,
                "step_pt": 2,
                "max_iterations": 20,
            }
        ],
        _ValvesStub(),
    )
    assert mutated == 1
    after = mod._edit_shape_font_pt_effective(shape)
    assert after <= 20.5, f"to_min expected ~20pt, got {after}"
    print(f"OK: endParaRPr effective + shrink to_min ({eff}pt -> {after}pt)")


def _test_set_font_pt(mod) -> None:
    data, sid = _build_spc_fixture()[0:2]
    prs = Presentation(BytesIO(data))
    mod._apply_edit_operations(
        prs,
        [{"op": "set_font_pt", "slide": 0, "shape_id": sid, "font_pt": 14}],
        _ValvesStub(),
    )
    shape = mod._shape_by_id_recursive(prs.slides[0], sid)
    assert mod._edit_shape_font_pt_effective(shape) == 14.0
    print("OK: set_font_pt (absolute)")


def _test_shrink_pe1b_to_min_fits(mod) -> None:
    data, sid = _build_spc_fixture()[0:2]
    prs = Presentation(BytesIO(data))
    shape = mod._shape_by_id_recursive(prs.slides[0], sid)
    before = mod._edit_shape_font_pt_effective(shape)
    assert mod._text_fits_shape_heuristic(shape)
    changed = mod._apply_edit_operations(
        prs,
        [
            {
                "op": "shrink_font",
                "slide": 0,
                "shape_id": sid,
                "shrink_mode": "to_min",
                "min_pt": 10,
                "step_pt": 2,
                "max_iterations": 20,
            }
        ],
        _ValvesStub(),
    )
    assert changed == 1
    after = mod._edit_shape_font_pt_effective(shape)
    assert after < before
    print(f"OK: shrink_font to_min when fit (PE1b) {before}pt -> {after}pt")


def main() -> None:
    global mod
    mod = _load_mod()
    _test_shrink_pe1(mod)
    _test_set_text_pe2(mod)
    _test_parse_and_apply_errors(mod)
    asyncio.run(_test_edit_tool_guards(mod))
    _test_marketing_reuse_post_edit(mod)
    _test_v2_ops(mod)
    _test_marketing_autofit_and_inspect(mod)
    _test_endpara_effective_and_to_min(mod)
    _test_set_font_pt(mod)
    _test_shrink_pe1b_to_min_fits(mod)
    print("edit_presentation_smoke: all checks passed")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"FAIL: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc
