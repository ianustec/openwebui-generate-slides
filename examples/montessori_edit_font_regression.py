"""Regression: Montessori deck font edits (fit vs to_min / set_font_pt).

    pip install -r examples/requirements-dev.txt
    python examples/montessori_edit_font_regression.py

Requires examples/output/indian_motessori-edit-simulation.pptx and .json
(generate locally or copy from chat export). Skips if missing.
"""
from __future__ import annotations

import importlib.util
import json
import sys
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from pptx import Presentation

BASE = Path(__file__).resolve().parent.parent
PPTX = BASE / "examples" / "output" / "indian_motessori-edit-simulation.pptx"
JSON = BASE / "examples" / "output" / "indian_motessori-edit-simulation.json"

# Title targets (PowerPoint slide number, shape_id) from inspect on generated deck
TITLE_TARGETS = [
    (2, 1100),
    (3, 1105),
    (5, 1462),
    (11, 1862),
    (14, 2252),
    (19, 2624),
]


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(mod)
    return mod


class _ValvesStub:
    presentation_edit_min_font_pt = 10.0
    presentation_edit_shrink_step_pt = 2.0


def _resolve_targets(mod, prs) -> list[tuple[int, int]]:
    found: list[tuple[int, int]] = []
    for slide_ui, sid in TITLE_TARGETS:
        try:
            internal = mod._edit_slide_json_to_internal(slide_ui)
        except ValueError:
            continue
        if internal < 0 or internal >= len(prs.slides):
            continue
        sh = mod._shape_by_id_recursive(prs.slides[internal], sid)
        if sh is not None and getattr(sh, "has_text_frame", False):
            found.append((slide_ui, sid))
    return found


def _font_on_slide(mod, prs, slide_ui: int, shape_id: int) -> float | None:
    internal = mod._edit_slide_json_to_internal(slide_ui)
    sh = mod._shape_by_id_recursive(prs.slides[internal], shape_id)
    if sh is None or not getattr(sh, "has_text_frame", False):
        return None
    return mod._edit_shape_font_pt_effective(sh)


def _count_title_font_drops(
    mod, prs_before, prs_after, targets: list[tuple[int, int]], threshold_pt: float
) -> int:
    n = 0
    for slide_ui, sid in targets:
        b = _font_on_slide(mod, prs_before, slide_ui, sid)
        a = _font_on_slide(mod, prs_after, slide_ui, sid)
        if b is None or a is None:
            continue
        if a < b - 0.5 and a <= threshold_pt + 1.0:
            n += 1
    return n


def _ops_from_json(
    mod,
    prs,
    spec: dict,
    targets: list[tuple[int, int]],
    *,
    mode: str | None,
    use_set_font: bool,
) -> list:
    v = _ValvesStub()
    raw_ops = spec.get("operations") or []
    if use_set_font:
        out = []
        for slide_ui, sid in targets:
            out.append(
                mod._normalize_edit_op(
                    {
                        "op": "set_font_pt",
                        "slide": slide_ui,
                        "shape_id": sid,
                        "font_pt": 22,
                    },
                    len(out),
                    v,
                )
            )
        return out
    out = []
    idx = 0
    for op in raw_ops:
        o = deepcopy(op)
        if mode and o.get("op") in ("shrink_font", "fit_text"):
            o["mode"] = mode
        try:
            norm = mod._normalize_edit_op(o, idx, v)
            mod._resolve_edit_shape(
                prs,
                norm["slide"],
                norm.get("shape_id") or norm.get("from_shape_id"),
            )
            out.append(norm)
            idx += 1
        except ValueError:
            continue
    return out


def main() -> None:
    if not PPTX.is_file():
        print(f"SKIP: missing {PPTX}")
        return
    mod = _load_mod()
    raw = PPTX.read_bytes()
    spec = json.loads(JSON.read_text()) if JSON.is_file() else {"operations": []}

    prs_before = Presentation(BytesIO(raw))
    targets = _resolve_targets(mod, prs_before)
    if len(targets) < 3:
        print(
            f"SKIP: only {len(targets)} title targets resolved in {PPTX.name} "
            f"(need >=3)"
        )
        return

    # Scenario A: default fit shrink on simulation JSON — expect few/no title changes
    prs_a = Presentation(BytesIO(raw))
    ops_a = _ops_from_json(
        mod, prs_a, spec, targets, mode=None, use_set_font=False
    )
    if not ops_a:
        print("SKIP: no simulation ops resolved against this deck")
        return
    mutated_a = mod._apply_edit_operations(prs_a, ops_a, _ValvesStub())
    drops_a = _count_title_font_drops(mod, prs_before, prs_a, targets, 28.0)
    print(f"Scenario A (shrink fit): ops={len(ops_a)} mutated={mutated_a} title_drops={drops_a}")
    assert drops_a <= 2, (
        f"fit shrink should change at most ~1–2 titles, got {drops_a} "
        "(overflow-only behavior)"
    )

    # Scenario B: set_font_pt on known title shapes
    prs_b = Presentation(BytesIO(raw))
    ops_b = _ops_from_json(
        mod, prs_b, spec, targets, mode=None, use_set_font=True
    )
    mutated_b = mod._apply_edit_operations(prs_b, ops_b, _ValvesStub())
    drops_b = _count_title_font_drops(mod, prs_before, prs_b, targets, 24.0)
    print(f"Scenario B (set_font_pt 22): ops={len(ops_b)} mutated={mutated_b} title_drops={drops_b}")
    assert mutated_b >= max(1, len(targets) // 2)
    assert drops_b >= max(2, len(targets) - 2), (
        f"expected most titles at 22pt, got {drops_b} drops"
    )

    # Scenario C: same JSON with to_min on shrink ops
    prs_c = Presentation(BytesIO(raw))
    ops_c = _ops_from_json(
        mod, prs_c, spec, targets, mode="to_min", use_set_font=False
    )
    mutated_c = mod._apply_edit_operations(prs_c, ops_c, _ValvesStub())
    drops_c = _count_title_font_drops(mod, prs_before, prs_c, targets, 12.0)
    print(f"Scenario C (shrink to_min): ops={len(ops_c)} mutated={mutated_c} title_drops={drops_c}")
    assert mutated_c >= 1
    assert drops_c >= drops_a, "to_min should change at least as many titles as fit"

    print("montessori_edit_font_regression: all checks passed")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"FAIL: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc
