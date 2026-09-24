"""Fase 4: title_bullets in template mode — bullets inside reference safe zone.

    pip install -r examples/requirements-dev.txt
    python examples/template_mode_title_bullets.py
"""
from __future__ import annotations

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
FIXTURES = BASE / "examples" / "fixtures"
OUT_DIR = BASE / "examples" / "output"
TOL_IN = 0.08


def _load_mod():
    spec = importlib.util.spec_from_file_location(
        "generate_slides", BASE / "generate_slides.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _bbox_in(slide, mod):
    return mod._BBox(
        x=mod._emu_in(slide.left),
        y=mod._emu_in(slide.top),
        w=mod._emu_in(slide.width),
        h=mod._emu_in(slide.height),
    )


def _contains(outer: "mod._BBox", inner: "mod._BBox", tol: float) -> bool:
    return (
        inner.x >= outer.x - tol
        and inner.y >= outer.y - tol
        and inner.x + inner.w <= outer.x + outer.w + tol
        and inner.y + inner.h <= outer.y + outer.h + tol
    )


def main() -> None:
    mod = _load_mod()
    if not mod._HAS_PPTX:
        print("python-pptx not available", file=sys.stderr)
        sys.exit(2)

    from pptx import Presentation

    corners = FIXTURES / "template_corners_one_slide.pptx"
    if not corners.is_file():
        print(f"Missing {corners}", file=sys.stderr)
        sys.exit(1)

    source_bytes = corners.read_bytes()
    pack = mod._parse_reference_pptx(source_bytes)
    pack = mod._apply_template_edits(pack, None)
    sz = pack.slides[0].safe_zone
    assert sz is not None, "fixture must have computed safe_zone"

    spec = {
        "title": "Template mode test",
        "template_mapping": {"default": 0, "content": 0},
        "slides": [
            {
                "layout": "title_bullets",
                "title": "Key points",
                "bullets": [
                    "First item inside safe zone",
                    "Second item",
                    "Third item",
                ],
            }
        ],
    }

    data, n = mod.Tools()._build(
        spec, template_pack=pack, reference_bytes=source_bytes
    )
    assert n == 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / "template_mode_title_bullets.pptx"
    out_path.write_bytes(data)

    prs = Presentation(BytesIO(data))
    slide = prs.slides[0]
    pics = [sh for sh in slide.shapes if mod._shape_kind(sh) == "picture"]
    assert len(pics) >= 4, f"expected >=4 corner pictures, got {len(pics)}"

    bullet_shapes = []
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        text = (sh.text_frame.text or "").strip()
        if "First item inside safe zone" in text or "Second item" in text:
            bullet_shapes.append(sh)

    assert bullet_shapes, "no bullet textbox found on output slide"
    for sh in bullet_shapes:
        bb = _bbox_in(sh, mod)
        if not _contains(sz, bb, TOL_IN):
            print(
                f"FAIL: bullet bbox {bb.x:.2f},{bb.y:.2f},{bb.w:.2f},{bb.h:.2f} "
                f"outside safe_zone {sz.x:.2f},{sz.y:.2f},{sz.w:.2f},{sz.h:.2f}",
                file=sys.stderr,
            )
            sys.exit(1)

    print(f"OK: wrote {out_path} ({len(pics)} pictures, bullets in safe zone)")
    sys.exit(0)


if __name__ == "__main__":
    main()
